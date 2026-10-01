"""Builds dist/data/casco-market.js from the case-study data and the model config.

The game uses the dataset as the market: how many customers sit in each city × channel × vehicle age ×
persona × customer type cell, what they paid and how the competition priced them. Claims come from the
model config (frequency, severity, loss ratio and the coefficient tables), never from the claim columns.

The model is calibrated to the data the way the data was generated:
  · the frequency and severity coefficient products are normalized to the portfolio (constants read from
    the data's freq_coef_combined / sev_coef_combined columns);
  · the market premium coefficients are the data's own (COEFFICIENTS sheet);
  · the base severity is set so the model's reference premium (base frequency × severity / loss ratio ×
    premium coefficients × competitor index) reproduces the premiums in the data.

Usage:  python3 scripts/build-casco-market.py <data.xlsx> <model_config.xlsx>
        python3 scripts/build-casco-market.py                    (defaults to the two files in veri/)
"""
import io, json, sys, zipfile, collections
import openpyxl

args = sys.argv[1:] or ["veri/pricing_case_data.xlsx", "veri/insurers_league_model_config.xlsx"]
if args[0].endswith('.zip'):
    with zipfile.ZipFile(args[0]) as z:
        names = z.namelist()
        data_name = next(n for n in names if 'config' not in n.lower())
        config_name = next(n for n in names if 'config' in n.lower())
        data_wb = openpyxl.load_workbook(io.BytesIO(z.read(data_name)), read_only=True, data_only=True)
        config_wb = openpyxl.load_workbook(io.BytesIO(z.read(config_name)), data_only=True)
    src = args[0]
else:
    data_wb = openpyxl.load_workbook(args[0], read_only=True, data_only=True)
    config_wb = openpyxl.load_workbook(args[1], data_only=True)
    src = f'{args[0].split("/")[-1]} + {args[1].split("/")[-1]}'

# ——— Model config ———
def table(name):
    rows = [r for r in config_wb[name].iter_rows(values_only=True) if r and r[0] is not None]
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:]]

model = {r['V2 CASE STUDY MODEL']: r['Value'] for r in table('MODEL_CONFIG')}
dims = {}
for key, sheet in [('city', 'CITY'), ('channel', 'CHANNEL'), ('vehicle', 'VEHICLE'), ('persona', 'PERSONA'), ('type', 'CUSTOMER_TYPE')]:
    dims[key] = [{
        'id': r['Level'], 'freq': r['Freq Coef'], 'sev': r['Severity Coef'], 'prem': r['Premium Coef'],
        **({'expense': r['Expense Ratio']} if 'Expense Ratio' in r else {})
    } for r in table(sheet)]

# The data's own market premium coefficients win over the config's (the data's premiums were built with them).
DIM_OF = {'City': 'city', 'Channel': 'channel', 'Vehicle Segment': 'vehicle', 'Persona': 'persona', 'Customer Type': 'type'}
if 'COEFFICIENTS' in data_wb.sheetnames:
    crow = data_wb['COEFFICIENTS'].iter_rows(values_only=True)
    chead = next(crow)
    ci = {h: i for i, h in enumerate(chead) if h}
    for r in crow:
        if not r or r[0] not in DIM_OF or 'Premium Coef' not in ci or r[ci['Premium Coef']] is None:
            continue
        for lv in dims[DIM_OF[r[0]]]:
            if lv['id'] == r[ci['Level']]:
                lv['prem'] = round(float(r[ci['Premium Coef']]), 4)

# ——— Market profile from the rows ———
rows = data_wb['DATA'].iter_rows(values_only=True)
head = next(rows)
col = {h: i for i, h in enumerate(head) if h}
cells = collections.OrderedDict()
total = 0
lv_of = {k: {x['id']: x for x in v} for k, v in dims.items()}
order = ['city', 'channel', 'vehicle', 'persona', 'type']
norm = {'f': [0.0, 0], 's': [0.0, 0]}
paid = rel = 0.0
for r in rows:
    if not r or not r[col['customer_id']]:
        continue
    key = (r[col['city']], r[col['renewal_channel']], r[col['vehicle_segment']], r[col['persona']], r[col['customer_type']])
    for dim, level in zip(order, key):
        if level not in lv_of[dim]:
            raise SystemExit(f'Unknown {dim} level in data: {level}')
    c = cells.setdefault(key, [0, 0.0, 0.0])
    c[0] += 1
    c[1] += float(r[col['base_premium_prev_term']])
    c[2] += float(r[col['competitor_price_index']])
    total += 1
    raw_f = raw_s = raw_p = 1.0
    for dim, level in zip(order, key):
        raw_f *= lv_of[dim][level]['freq']; raw_s *= lv_of[dim][level]['sev']; raw_p *= lv_of[dim][level]['prem']
    if 'freq_coef_combined' in col:
        norm['f'][0] += float(r[col['freq_coef_combined']]) / raw_f; norm['f'][1] += 1
        norm['s'][0] += float(r[col['sev_coef_combined']]) / raw_s; norm['s'][1] += 1
    paid += float(r[col['final_premium_gross']])
    rel += raw_p * float(r[col['competitor_price_index']])
freq_norm = round(norm['f'][0] / norm['f'][1], 6) if norm['f'][1] else 1
sev_norm = round(norm['s'][0] / norm['s'][1], 6) if norm['s'][1] else 1
# Reference premium unit that reproduces the data's premiums, and the severity it implies.
unit = paid / rel
severity = round(unit * model['Base Loss Ratio'] / model['Base Frequency'], 2)

index = {k: {x['id']: i for i, x in enumerate(v)} for k, v in dims.items()}
out_cells = []
for key, (n, prem, comp) in cells.items():
    out_cells.append([*(index[d][lvl] for d, lvl in zip(order, key)), n, round(prem / n, 2), round(comp / n, 4)])
out_cells.sort()

market = {
    'source': src.split('/')[-1],
    'customers': total,
    'model': {'frequency': model['Base Frequency'], 'severity': severity, 'lossRatio': model['Base Loss Ratio'], 'gammaShape': model['Gamma Shape'],
              'freqNorm': freq_norm, 'sevNorm': sev_norm},
    'calibration': {'configSeverity': model['Base Severity EUR'], 'premiumUnit': round(unit, 2), 'dataPremium': round(paid / total, 2)},
    'dimensions': dims,
    'cellColumns': order + ['count', 'avgPremium', 'competitorIndex'],
    'cells': out_cells,
}
with open('dist/data/casco-market.js', 'w') as f:
    f.write('// Generated by scripts/build-casco-market.py from the case-study data. Do not edit by hand.\n')
    f.write('export const CASCO_MARKET = ')
    json.dump(market, f, ensure_ascii=False, separators=(',', ':'))
    f.write(';\n')
print(f'{total} customers in {len(out_cells)} cells → dist/data/casco-market.js')
print(f'freq norm {freq_norm} · sev norm {sev_norm} · premium unit €{unit:.2f} → base severity €{severity} (config said €{model["Base Severity EUR"]})')
