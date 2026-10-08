"""Builds oyun-varsayimlari.xlsx: every variable and assumption the game uses, one topic per sheet.

Values are read from the engine itself (dist/), so the book always matches what the game runs.
Run again whenever the rules or the data change:  python3 scripts/build-assumptions-book.py
"""
import json, subprocess, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DUMP = r"""
const e = await import('./dist/engine.js');
const c = await import('./dist/casco.js');
const g = await import('./dist/js/game.js');
const { CASCO_MARKET } = await import('./dist/data/casco-market.js');
const cfg = e.scenario('tr'), R = e.rulesOf(cfg), money = e.cascoMoney(cfg), policies = e.policiesOf(cfg);
const dims = e.DIMENSIONS;
const levels = Object.fromEntries(dims.map(d => [d, R.dimensions[d].map((lv, i) => ({ ...lv, name: e.levelName(d, i, 'tr') }))]));
const dimNames = Object.fromEntries(dims.map(d => [d, e.dimensionName(d, 'tr')]));
const act = e.actuarialCoefficients(R);
const cells = CASCO_MARKET.cells.map(cell => { const r = e.cellRisk(cell, R); return { cell, ...r }; });
const ref = e.referenceMarket(R, policies);
const scopes = Object.fromEntries(e.EVENT_SCOPES.map(sc => { if (sc.id === 'all') return [sc.id, 'Tüm pazar']; const [d, i] = sc.id.split(':'); return [sc.id, `${dimNames[d]}: ${e.levelName(d, Number(i), 'tr')}`]; }));
console.log(JSON.stringify({
  assumptions: cfg.assumptions, weights: cfg.weights, events: cfg.events, speed: cfg.speed, seed: cfg.seed, year: cfg.year,
  rules: { ...R, dimensions: undefined }, levels, dimNames, dims, act, money, policies, ref, scopes,
  customers: CASCO_MARKET.customers, source: CASCO_MARKET.source, calibration: CASCO_MARKET.calibration,
  cells, nominal: c.NOMINAL_TEAMS, quarterKeys: c.QUARTER_KEYS,
  audienceYear: e.campaignAudience(R, money), offerNames: Object.fromEntries(e.campaignRules(R).offers.map(o => [o.id, e.offerName(o.id, 'tr')])),
  digitalPolicies: CASCO_MARKET.cells.filter(x => x[1] === e.campaignRules(R).channel).reduce((a, x) => a + x[5], 0) / CASCO_MARKET.customers * policies,
  session: { minTeams: g.MIN_TEAMS, maxTeams: g.MAX_TEAMS, reviewMonths: g.STRATEGY_REVIEW_MONTHS, reviewSeconds: g.EXCEL_REVIEW_MS / 1000, finalDelay: g.FINAL_DELAY_MS / 1000, minutes: cfg.minutes }
}));
"""
D = json.loads(subprocess.run(['node', '--input-type=module', '-e', DUMP], capture_output=True, text=True, check=True).stdout)
R, A, M = D['rules'], D['assumptions'], D['money']
LV, DN, DIMS = D['levels'], D['dimNames'], D['dims']

INK, MUTED, LINE = '1E1E1E', '5B6470', 'C3C8CE'
HEAD = PatternFill('solid', fgColor='1E1E1E')
SECTION = PatternFill('solid', fgColor='FFE066')
NOTE = PatternFill('solid', fgColor='FFF6D6')
thin = Border(bottom=Side(style='thin', color=LINE))
EUR, EUR2, PCT, PCT1, N0, N2, N4 = '#,##0 "€"', '#,##0.00 "€"', '0%', '0.0%', '#,##0', '0.00', '0.0000'

wb = Workbook()
wb.remove(wb.active)


def sheet(title, headers, widths, intro=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    r = 1
    ws.cell(r, 1, title).font = Font(bold=True, size=16, color=INK)
    r += 1
    if intro:
        ws.cell(r, 1, intro).font = Font(size=11, color=MUTED, italic=True)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        ws.row_dimensions[r].height = max(30, 15 * (len(intro) // 120 + 1))
        r += 1
    r += 1
    for i, h in enumerate(headers, 1):
        c = ws.cell(r, i, h)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = ws.cell(r + 1, 1)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws._next = r + 1
    return ws


def row(ws, values, fmts=None, fill=None, bold=False):
    r = ws._next
    for i, v in enumerate(values, 1):
        c = ws.cell(r, i, v)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = thin
        if fmts and i - 1 < len(fmts) and fmts[i - 1]:
            c.number_format = fmts[i - 1]
        if fill:
            c.fill = fill
        if bold:
            c.font = Font(bold=True)
    ws._next = r + 1
    return r


def section(ws, text, span):
    r = ws._next
    c = ws.cell(r, 1, text)
    c.font, c.fill = Font(bold=True, color=INK), SECTION
    for i in range(2, span + 1):
        ws.cell(r, i).fill = SECTION
    ws._next = r + 1


cells = D['cells']
tr = lambda v, d: f'{v:,.{d}f}'.replace(',', ' ').replace('.', ',').replace(' ', '.')
n_total = sum(x['cell'][5] for x in cells)
port_freq = sum(x['cell'][5] * x['freq'] for x in cells) / n_total
port_cost = sum(x['cell'][5] * x['cost'] for x in cells) / n_total
port_ref = sum(x['cell'][5] * x['reference'] for x in cells) / n_total
port_sev = port_cost / port_freq

# ——— 1. Oku beni ———
ws = sheet('Oku beni', ['Sayfa', 'İçerik'], [28, 110],
           'Insurers League kasko yarışının kullandığı bütün değişkenler ve varsayımlar. Değerler oyunun motorundan okunmuştur; oyun ne çalıştırıyorsa burada o yazar.')
for name, text in [
    ('Pazar ve para', 'Pazar büyüklüğü, sermaye, karar bütçesi, sabit gider, reasürans bedeli ve rekabet ayarları.'),
    ('Hasar modeli', 'Frekans, hasar tutarı, hasar/prim oranı, dağılım parametreleri ve veriye kalibrasyon.'),
    ('Segment katsayıları', 'Beş boyutun (il, kanal, araç yaşı, persona, müşteri tipi) her seviyesi için frekans, şiddet, piyasa prim katsayısı ve gider oranı.'),
    ('Müşteri davranışı', 'Her personanın fiyata, hizmete ve kanala ne kadar duyarlı olduğu. Takımlara verilmez; veriden ve yarıştan çıkarırlar.'),
    ('Pazarlama ve hizmet', 'Pazarlama harcamasının görünürlüğe, hasar operasyonu bütçesinin kapasiteye ve memnuniyete (NPS) dönüşümü.'),
    ('Dijital kampanya', 'Marketing_Input dosyasından gelen müşteri kazanma kampanyası: medya maliyeti, hedef kitle, hediyeler ve huni.'),
    ('Reasürans', 'Tek kota paylı anlaşmanın devir oranı, komisyonu ve bedeli.'),
    ('Takım kararları', 'Takımların teslim ettiği her karar, sınırları ve çeyrek molalarında değişip değişemeyeceği.'),
    ('Olay takvimi', 'Sezon boyunca pazara düşen olaylar: ay, süre, kapsam, hasar ve talep çarpanı.'),
    ('Oturum ayarları', 'Takım sayısı, çeyrek molaları, süreler, tohum (seed) ve oturumla ilgili diğer ayarlar.'),
    ('Ödüller ve göstergeler', 'Üç kupa, sahnede gösterilen göstergeler ve arka plandaki birleşik puanın tanımı.'),
    ('Hesap adımları', 'Bir ayın motor içinde nasıl hesaplandığı, sırasıyla.'),
    ('Pazar hücreleri', f'Örneklemdeki {len(cells)} segment hücresi: müşteri adedi, yıllık poliçe, referans prim, beklenen hasar ve hasar/prim oranı.'),
    ('Veri kullanımı', 'Vaka verisinin her sütunu: oyunda kullanılıyor mu, nasıl.'),
    ('Denetim bulguları', 'Veri, girdi dosyaları ve oyun mantığı arasındaki tutarsızlıklar; durumu ve öneri.'),
]:
    row(ws, [name, text])
ws._next += 1
row(ws, ['Kaynak veri', D['source']], fill=NOTE)
row(ws, ['Örneklem', f"{D['customers']:,} müşteri → {len(cells)} hücre".replace(',', '.')], fill=NOTE)
row(ws, ['Üretildiği tarih', datetime.date.today().strftime('%d.%m.%Y')], fill=NOTE)
row(ws, ['Yeniden üretmek için', 'python3 scripts/build-assumptions-book.py'], fill=NOTE)
row(ws, ['Not', 'Bu değerler yeni bir oyunun varsayılanlarıdır. Moderatör çoğunu Kural stüdyosundan oyun başlamadan değiştirebilir; açık bir oda kendi ayarlarıyla devam eder.'], fill=NOTE)

# ——— 2. Pazar ve para ———
ws = sheet('Pazar ve para', ['Değişken', 'Değer', 'Birim', 'En az', 'En çok', 'Ne yapar', 'Nereden gelir'], [30, 16, 10, 12, 16, 70, 44],
           'Para değerleri, referans pazarın bir takıma düşen adil diliminin (pazar ÷ 6) payı olarak türetilir; pazar büyüklüğü değişince orantılı ölçeklenir.')
section(ws, 'Pazar', 7)
row(ws, ['Yıllık pazar büyüklüğü', A['policies']['value'], 'poliçe', A['policies']['min'], A['policies']['max'],
         'Tüm pazarın bir yılda aldığı poliçe. Her hücre örneklemdeki payı kadar poliçe alır; profil örneklemle birebir aynıdır.', 'Moderatör (Kural stüdyosu)'], [None, N0, None, N0, N0])
row(ws, ['Aylık müşteri havuzu (ortalama)', round(A['policies']['value'] / 12), 'poliçe', None, None, 'Her ay pazara çıkan müşteri; mevsimsellik, olaylar ve küçük bir rastgelelikle oynar.', 'Yıllık ÷ 12'], [None, N0])
row(ws, ['Referans pazar primi (yıllık)', D['ref']['gwp'], '€', None, None, 'Bütün pazar piyasa fiyatından alsaydı yazılacak brüt prim.', 'Σ hücre poliçesi × referans prim'], [None, EUR])
row(ws, ['Referans pazar hasarı (yıllık, beklenen)', D['ref']['claims'], '€', None, None, 'Aynı pazarın beklenen hasar maliyeti.', 'Σ hücre poliçesi × frekans × şiddet'], [None, EUR])
row(ws, ['Adil dilim (bir takım)', M['slice'], '€', None, None, 'Referans pazarın 6 takıma eşit bölünmüş hali. Para varsayılanları bunun payıdır.', 'Referans prim ÷ 6'], [None, EUR])
row(ws, ['Dış seçenek ağırlığı', R['market']['outside'], 'katsayı', None, None, 'Takımlardan hiçbirini seçmeyen müşterinin (piyasanın geri kalanı) seçim ağırlığı. Büyüdükçe takımların toplam payı küçülür.', 'Kural stüdyosu'], [None, N2])
row(ws, ['Mevsimsellik', R['market']['seasonality'], 'oran', None, None, 'Aylık talep 1 + bu oran × sin((ay+1)/1,9) ile dalgalanır.', 'Kural stüdyosu'], [None, PCT])
section(ws, 'Para (her takıma aynı)', 7)
for key, share, extra in [('capital', 0.30, 'Her şirketin başlangıç özkaynağı. Özkaynak kuralı kapalı: eksiye düşen takım yine kupa alabilir.'),
                          ('budget', 0.08, 'Pazarlama + hasar operasyonu + (seçilirse) reasürans bedeli; yıla eşit yayılır.'),
                          ('fixedCost', 0.02, 'Her şirketin taşıdığı yıllık genel gider; aylık 1/12 düşer.'),
                          ('reinsuranceFee', 0.01, 'Kota paylı anlaşmanın sabit bedeli; karar bütçesinden ödenir.')]:
    a = A[key]
    row(ws, [a['label'], a['value'], '€', a['min'], a['max'], extra, f'Adil dilimin %{share * 100:.0f}’i (Kural stüdyosundan değişir)'], [None, EUR, None, EUR, EUR])
row(ws, ['Pazarlama ölçeği', M['marketingScale'], '€', None, None, 'Pazarlama getirisinin azalmaya başladığı harcama düzeyi (görünürlük formülünde payda).', f"Adil dilim × {tr(R['marketing']['scale'], 3)}"], [None, EUR])

# ——— 3. Hasar modeli ———
ws = sheet('Hasar modeli', ['Değişken', 'Değer', 'Birim', 'Ne yapar', 'Kaynak'], [34, 16, 10, 74, 50],
           'Hasarlar yalnızca model konfigürasyonundan üretilir; verideki hasar sütunları okunmaz. Model, veriyi üreten mantıkla aynı şekilde normalize ve kalibre edilmiştir.')
section(ws, 'Model parametreleri', 5)
m = R['model']
row(ws, ['Temel frekans', m['frequency'], 'adet/yıl', 'Poliçe başına yıllık beklenen hasar adedi (portföy ortalaması).', 'Model konfigürasyonu · Base Frequency'], [None, N4])
row(ws, ['Temel hasar tutarı (şiddet)', m['severity'], '€', 'Bir hasarın baz maliyeti. Segment şiddet katsayılarıyla çarpılır.', f"Veriye kalibre edildi (konfigürasyonda €{D['calibration']['configSeverity']})"], [None, EUR2])
row(ws, ['Piyasa hasar/prim oranı (baz)', m['lossRatio'], 'oran', 'Piyasa referans priminin baz birimini belirler: frekans × şiddet ÷ bu oran.', 'Model konfigürasyonu · Base Loss Ratio'], [None, PCT])
row(ws, ['Gamma şekil parametresi', m['gammaShape'], '', 'Hasar tutarının sağa çarpıklığı. Küçük değer = daha oynak hasar.', 'Model konfigürasyonu · Gamma Shape'], [None, N2])
row(ws, ['Frekans normalizasyonu', m['freqNorm'], 'çarpan', 'Beş frekans katsayısının çarpımını portföy ortalamasına ölçekler (veride freq_coef_combined ile aynı).', 'Veriden ölçüldü'], [None, N4])
row(ws, ['Şiddet normalizasyonu', m['sevNorm'], 'çarpan', 'Beş şiddet katsayısının çarpımını portföy ortalamasına ölçekler (veride sev_coef_combined ile aynı).', 'Veriden ölçüldü'], [None, N4])
section(ws, 'Kalibrasyon', 5)
row(ws, ['Verideki ortalama prim', D['calibration']['dataPremium'], '€', 'Örneklemdeki final_premium_gross ortalaması.', 'Veri'], [None, EUR2])
row(ws, ['Referans prim birimi', D['calibration']['premiumUnit'], '€', 'Prim katsayısı ve rakip endeksi 1 olan müşterinin piyasa fiyatı; verideki primleri yeniden üretir.', 'Σ prim ÷ Σ(prim katsayısı × rakip endeksi)'], [None, EUR2])
row(ws, ['Konfigürasyondaki şiddet', D['calibration']['configSeverity'], '€', 'Bu değerle verideki primlere karşı hasar/prim oranı yaklaşık %4,5 çıkıyordu; kullanılmıyor.', 'Model konfigürasyonu'], [None, EUR2])
section(ws, 'Sonuç: portföy ortalamaları (referans fiyatta)', 5)
row(ws, ['Hasar frekansı', port_freq, 'adet/yıl', 'Poliçe ağırlıklı ortalama.', 'Hesaplanan'], [None, PCT1])
row(ws, ['Ortalama hasar tutarı', port_sev, '€', 'Frekans ağırlıklı ortalama şiddet.', 'Hesaplanan'], [None, EUR])
row(ws, ['Poliçe başına beklenen hasar', port_cost, '€', 'Frekans × şiddet.', 'Hesaplanan'], [None, EUR2])
row(ws, ['Ortalama referans prim', port_ref, '€', 'Piyasa fiyatı (veriyle aynı).', 'Hesaplanan'], [None, EUR2])
row(ws, ['Piyasa hasar/prim oranı', port_cost / port_ref, 'oran', 'Referans fiyatta beklenen brüt hasar/prim; verinin kendi beklenen oranıyla tutarlı.', 'Hesaplanan'], [None, PCT1])
section(ws, 'Hasar çekimi', 5)
row(ws, ['Hasar adedi', 'Poisson (normal yaklaşımı)', '', 'Takımın o ay sattığı poliçelerin beklenen hasar adedi etrafında; takım ve ay için sabit şans tablosundan.', 'Motor'])
row(ws, ['Hasar tutarı', 'Gamma', '', 'Adet × ortalama şiddet; şekil parametresi yukarıda. Aynı tohum (seed) aynı sezonu üretir.', 'Motor'])

# ——— 4. Segment katsayıları ———
ws = sheet('Segment katsayıları', ['Boyut', 'Seviye', 'Veride', 'Örneklem payı', 'Frekans kat.', 'Şiddet kat.', 'Piyasa prim kat.', 'Kanal gider oranı',
                                    'Beklenen hasar/prim (piyasa fiyatında)', 'Risk bazlı katsayı (cevap anahtarı)'],
           [16, 20, 18, 12, 11, 11, 13, 12, 18, 18],
           'Piyasa prim katsayıları verideki COEFFICIENTS tablosundan alınır. Son sütun moderatör içindir: segmentin riskine göre adil katsayı (0,05 adıma yuvarlı). Takımlara verilmez.')
for k, dim in enumerate(DIMS):
    section(ws, DN[dim], 10)
    for i, lv in enumerate(LV[dim]):
        sub = [x for x in cells if x['cell'][k] == i]
        share = sum(x['cell'][5] for x in sub) / n_total
        lr = sum(x['cell'][5] * x['cost'] for x in sub) / sum(x['cell'][5] * x['reference'] for x in sub)
        row(ws, [DN[dim], lv['name'], lv['id'], share, lv['freq'], lv['sev'], lv['prem'], lv.get('expense'), lr, D['act'][dim][i]],
            [None, None, None, PCT1, N4, N4, N4, PCT, PCT1, N2])

# ——— 5. Müşteri davranışı ———
ws = sheet('Müşteri davranışı', ['Persona', 'Fiyat duyarlılığı (β)', 'Hizmet ağırlığı', 'Kanal sadakati', 'Ne anlama gelir'], [20, 18, 16, 16, 80],
           'Bir müşterinin bir takımı seçme ağırlığı = (teklif ÷ referans prim)^(−β) × kanal görünürlüğü^sadakat × e^(hizmet ağırlığı × itibar × (hizmet skoru − 98) ÷ 40). Pay = ağırlık ÷ (dış seçenek + Σ ağırlık).')
notes = ['Bankasıyla çalışır; kanala bağlı, fiyata orta duyarlı, hizmete az bakar.',
         'Hasar yaşamış; fiyata duyarlı ama hizmeti çok önemser.',
         'En büyük grup; fiyatı en çok önemser, hizmete pek bakmaz.',
         'Kaliteye bakar; fiyata en az duyarlı, hizmeti önemser.']
for i, b in enumerate(R['behavior']):
    row(ws, [LV['persona'][i]['name'], b['price'], b['service'], b['channel'], notes[i]], [None, N2, N2, N2])
ws._next += 1
row(ws, ['Ticari müşteri fiyat çarpanı', R['commercialPrice'], None, None, 'Ticari müşterinin fiyat duyarlılığı persona β’sının bu katıdır (daha az duyarlı).'], [None, N2], fill=NOTE)
row(ws, ['Hizmet itibarı', R['service']['reputation'], None, None, 'Hizmet skorunun seçime etkisini topluca büyütür ya da küçültür.'], [None, N2], fill=NOTE)

# ——— 6. Pazarlama ve hizmet ———
ws = sheet('Pazarlama ve hizmet', ['Değişken', 'Değer', 'Birim', 'Ne yapar'], [34, 14, 10, 90])
section(ws, 'Pazarlama', 4)
mk, sv = R['marketing'], R['service']
row(ws, ['Taban görünürlük', mk['presence'], 'katsayı', 'Hiç pazarlama yapmayan takımın bir kanaldaki görünürlüğü.'], [None, N2])
row(ws, ['Pazarlama gücü', mk['strength'], 'katsayı', 'Görünürlük = taban + güç × ln(1 + kanal harcaması ÷ pazarlama ölçeği). Getiri azalarak artar.'], [None, N2])
row(ws, ['Pazarlama ölçeği', mk['scale'], '× adil dilim', f"Payda; varsayılan oyunda €{M['marketingScale']:,.0f}.".replace(',', '.')], [None, N4])
row(ws, ['Görünürlüğe giden pazarlama', 'pazarlama × (1 − kampanya %)', '', 'Kampanyaya ayrılan pay kanal görünürlüğü almaz; ikisi arasında takım seçim yapar.'])
row(ws, ['Kanal başına en az odak', mk['minFocus'], '%', 'Pazarlama odağı dört satış kanalına dağıtılır; her biri en az bu kadar, toplam %100.'], [None, N0])
section(ws, 'Hasar operasyonu ve memnuniyet', 4)
row(ws, ['Dosya başına işlem maliyeti', sv['handling'], '× şiddet', 'Aylık kapasite (dosya) = hasar operasyonu bütçesi ÷ 12 ÷ (bu oran × temel şiddet).'], [None, N2])
row(ws, ['Kullanım eşiği', sv['threshold'], 'oran', 'Gelen dosya ÷ kapasite bu eşiği aşınca hizmet skoru düşmeye başlar.'], [None, PCT])
row(ws, ['Düşüş eğimi', sv['slope'], 'puan', 'Eşiğin üstündeki her 1,0 kullanım için hizmet skorundan düşen puan.'], [None, N0])
row(ws, ['En düşük hizmet skoru', sv['floor'], 'puan', 'Hizmet skoru bunun altına inmez (NPS −40).'], [None, N0])
row(ws, ['En yüksek hizmet skoru', sv['max'], 'puan', 'Kapasite yeterliyse hizmet skoru (NPS +96).'], [None, N0])
row(ws, ['Hasar sızıntısı', sv['leakage'], 'oran', 'Kapasite aşılınca (kullanım > 1) hasarlar 1 + bu oran × (kullanım − 1) kadar büyür; en çok 3 birim aşım sayılır.'], [None, PCT])
row(ws, ['NPS', '2 × ort. hizmet − 100', '', 'Yılın ortalama hizmet skorunun −100…+100 ölçeğine çevrilmiş hali.'])

# ——— Dijital kampanya ———
K = R['campaign']
ws = sheet('Dijital kampanya', ['Değişken', 'Değer', 'Birim', 'Ne yapar'], [36, 16, 12, 92],
           'Marketing_Input.xlsx dosyasındaki kampanya: medya gösterim alır, ulaşılan kişi ilgi × tıklama × hit ile müşteriye dönüşür, her müşteri bir hediye alır. Vakadaki euro başına kişi korunarak pazarımıza ölçeklenmiştir.')
section(ws, 'Parametreler', 4)
row(ws, ['Medya maliyeti (CPM)', K['cpm'], '€ / 1.000', 'Bin gösterimin maliyeti.'], [None, N2])
row(ws, ['Dijital kullanıcı (vaka)', K['digitalUsers'], 'kişi', '18–55 yaş dijital kullanıcı (vaka metni).'], [None, N0])
row(ws, ['Hedef kitle payı', K['targetShare'], 'oran', '“Joyful Disregarders” (vaka metni).'], [None, PCT])
row(ws, ['Vakadaki kampanya bütçesi', K['referenceBudget'], '€', 'Bu bütçeye düşen kitle = kullanıcı × hedef payı. Euro başına kişi korunur.'], [None, EUR])
row(ws, ['Takım-yılı başına hedef kitle', D['audienceYear'], 'kişi', 'kullanıcı × hedef payı × karar bütçesi ÷ vakadaki bütçe. Tüm takımlar toplam kitleyi paylaşır.'], [None, N0])
row(ws, ['Aylık kitle (6 takımla)', D['audienceYear'] * 6 / 12, 'kişi', 'Takım sayısıyla büyür; ayrıca mevsimsellik ve kapsamdaki olayların talep çarpanıyla oynar.'], [None, N0])
row(ws, ['Frekans eşiği', K['frequency'], 'görüntüleme', 'Ortalama görüntüleme bunu aşarsa…'], [None, N0])
row(ws, ['Hit artışı', K['frequencyBonus'], 'oran', '…tüm kampanyaların hit oranı bu kadar artar.'], [None, PCT])
row(ws, ['Fiyat etkisi tavanı', K['priceCap'], 'kat', 'Piyasadan ucuz fiyat dönüşümü artırır, en fazla bu kat. Pahalı fiyat azaltır.'], [None, N2])
row(ws, ['Kampanya kanalı', 'Dijital', '', 'Kampanya müşterileri dijital kanal hücrelerine, örneklemdeki paylarıyla girer.'])
row(ws, ['Organik dijital pazar (yıllık)', D['digitalPolicies'], 'poliçe', 'Karşılaştırma için: kampanya müşterileri bunun üstüne yeni müşteri olarak eklenir.'], [None, N0])
section(ws, 'Hediyeler (vaka)', 4)
for o in K['offers']:
    conv = o['interest'] * o['click'] * o['hit']
    row(ws, [D['offerNames'][o['id']], f"ilgi %{o['interest'] * 100:.0f} · tıklama %{o['click'] * 100:.0f} · hit %{o['hit'] * 100:.0f}", f"€{o['cost']}",
             f"1.000 erişimde {conv * 1000:.1f} müşteri. Müşteri başı maliyet ≈ medya €{K['cpm'] / 1000 / conv:.2f} + hediye €{o['cost']} (hedef kitle doymadıysa)."])
section(ws, 'Kurallar', 4)
for a, b in [('Gösterim', 'medya bütçesi ÷ CPM × 1.000'), ('Erişim', 'gösterim × min(1, kitle ÷ tüm takımların gösterimi)'),
             ('Frekans', 'max(1, tüm gösterimler ÷ kitle) — tüm takımlar için aynı'), ('Müşteri', 'erişim × ilgi × tıklama × hit × fiyat etkisi'),
             ('Hediye sınırı', 'kazanılan müşteri ≤ hediye bütçesi ÷ hediye maliyeti; kullanılmayan hediye parası geri dönmez'),
             ('Muhasebe', 'Kampanya müşterileri takımın kendi fiyatıyla yeni poliçe olur; hasarı, kanal gideri ve hasar operasyonu yükü normal işler.')]:
    row(ws, [a, b, '', ''])

# ——— 7. Reasürans ———
ws = sheet('Reasürans', ['Değişken', 'Değer', 'Birim', 'Ne yapar'], [30, 14, 10, 90],
           'Tek anlaşma tipi: kota paylı. Takım yıl başında açar ya da açmaz; çeyrek molalarında değişmez.')
row(ws, ['Devir oranı', R['reinsurance']['share'], 'oran', 'Brüt primin ve hasarın bu kadarı reasüröre geçer.'], [None, PCT])
row(ws, ['Devir komisyonu', R['reinsurance']['commission'], 'oran', 'Devredilen primin bu kadarı takıma komisyon olarak geri döner (gideri azaltır).'], [None, PCT])
row(ws, ['Anlaşma bedeli', A['reinsuranceFee']['value'], '€', 'Sabit bedel; açılırsa karar bütçesinden düşer.'], [None, EUR])

# ——— 8. Takım kararları ———
ws = sheet('Takım kararları', ['Karar', 'Birim', 'Kural', 'Çeyrekte değişir mi', 'Varsayılan (dosya gelmezse)'], [30, 10, 70, 16, 34],
           'Takımlar bu kararları Excel şablonuyla teslim eder (kendi cihazlarından ya da moderatör yükler). Çeyrek molasında değişenler bir sonraki aydan itibaren geçerlidir.')
q = set(D['quarterKeys'])
yes = lambda k: 'Evet' if k in q else 'Hayır'
row(ws, ['Baz prim', '€', 'Sıfırdan büyük. Bir müşteriye teklif = baz prim × il × kanal × araç yaşı × persona × müşteri tipi katsayısı.', yes('basePremium'),
         f"€{tr(m['frequency'] * m['severity'] / m['lossRatio'], 2)} (frekans × şiddet ÷ hasar/prim)"])
row(ws, ['Segment katsayıları (19 adet)', 'çarpan', f"Her biri {tr(R['coef']['min'], 2)}–{tr(R['coef']['max'], 2)} arasında, {tr(R['coef']['step'], 2)} adımla (örn. 0,95 · 1,00 · 1,05).", yes('coef'), 'Hepsi 1,00'])
row(ws, ['Pazarlama bütçesi', '€', 'Sıfır ya da pozitif. Pazarlama + hasar operasyonu + reasürans bedeli ≤ karar bütçesi.', yes('marketing'), 'Bütçenin %50’si'])
row(ws, ['Pazarlama odağı (4 kanal)', '%', f"Acente, banka, dijital, broker; her biri en az %{mk['minFocus']}, toplam %100.", yes('channelFocus'), 'Her kanala %25'])
row(ws, ['Hasar operasyonu bütçesi', '€', 'Sıfır ya da pozitif; kapasiteyi ve dolayısıyla NPS’i belirler.', yes('claimsOps'), 'Bütçenin %40’ı'])
row(ws, ['Kampanya payı', '% pazarlamanın', '0–100, 5’er adım. Pazarlamanın bu kısmı kampanyaya, kalanı kanal görünürlüğüne gider.', yes('campaign'), '%20'])
row(ws, ['Medya payı', '% kampanyanın', '0–100, 5’er adım. Kalanı hediye bütçesidir.', yes('mediaShare'), '%50'])
row(ws, ['Kampanya hediyesi', 'liste', 'Konser/etkinlik indirimi, restoran kartı, kahve kartı, spor salonu indirimi; biri seçilir.', yes('offer'), 'Konser / etkinlik indirimi'])
row(ws, ['Reasürans (kota paylı)', 'açık/kapalı', 'Açılırsa anlaşma bedeli bütçeden düşer.', 'Hayır', 'Kapalı'])

# ——— 9. Olay takvimi ———
ws = sheet('Olay takvimi', ['Başladığı ay', 'Süre (ay)', 'Kapsam', 'Hasar çarpanı', 'Talep çarpanı', 'Başlık', 'Açıklama'], [12, 10, 26, 12, 12, 30, 80],
           'Olay, başladığı aydan itibaren süresi boyunca kapsamındaki hücrelere uygulanır. Hasar çarpanı o ay yazılan poliçelerin hasarını, talep çarpanı pazara çıkan müşteri sayısını değiştirir.')
for ev in D['events']:
    row(ws, [ev['month'] + 1, ev['duration'], D['scopes'].get(ev['scope'], ev['scope']), ev.get('cost', 1), ev.get('demand', 1), ev['title'], ev['description']], [N0, N0, None, N2, N2])

# ——— 10. Oturum ayarları ———
ws = sheet('Oturum ayarları', ['Ayar', 'Değer', 'Açıklama'], [32, 18, 90])
s = D['session']
row(ws, ['Sezon yılı', D['year'], 'Ekranda görünen yıl.'])
row(ws, ['Takım sayısı', f"{s['minTeams']}–{s['maxTeams']}", 'Yarış en az bu kadar takımla başlar.'])
row(ws, ['Karar süresi', f"{s['minutes']} dakika", 'Kararlar aşamasında geri sayım. Süre dolunca takımlar kendi cihazından yükleyemez; moderatör yükleyebilir ya da 5 dakika ekleyebilir.'])
row(ws, ['Çeyrek molaları', ', '.join(f"{x + 1}. ay sonu" for x in s['reviewMonths']), 'Mola sırasında takımlar güncel planını indirip yeni dosya yükleyebilir.'])
row(ws, ['Mola süresi', f"{s['reviewSeconds'] // 60} dakika", 'Süre dolunca ya da moderatör kapatınca yarış devam eder; dosya gelmeyen takımın planı aynen sürer.'])
row(ws, ['Ay başına süre', f"{D['speed']} sn", 'Moderatör 3–12 sn arasında değiştirebilir.'])
row(ws, ['Rastgelelik tohumu (seed)', D['seed'], 'Aynı kararlar ve aynı tohum aynı sezonu birebir üretir.'])
row(ws, ['Özkaynak kuralı', 'Kapalı' if not R['capitalRule'] else 'Açık', 'Açıksa özkaynağı sıfırın altına düşen takım kupa alamaz.'])
row(ws, ['Bilgi turu (quiz)', 'Kapalı', 'Takımlar cihazdan oynamadığı için kapalı.'])
row(ws, ['Muhasebe esası', 'Yazım yılı', 'Bir ayda satılan poliçe tam yıllık primini ve nihai hasarını o ay kaydeder.'])

# ——— 11. Ödüller ve göstergeler ———
ws = sheet('Ödüller ve göstergeler', ['Ad', 'Tanım', 'Eşitlikte'], [28, 90, 30],
           'Finalde üç ayrı kupa verilir; her biri kendi podyumuyla. Sahnede sıralama brüt prime göredir.')
section(ws, 'Kupalar', 3)
row(ws, ['Kârlılık', 'Yıl sonu teknik kâr = net prim − net hasar − giderler (kanal gideri + sabit gider + karar bütçesi − reasürans komisyonu).', '—'])
row(ws, ['Pazar payı', 'Yıl sonu brüt prim payı = takımın brüt primi ÷ takımların toplam brüt primi.', 'Brüt prim'])
row(ws, ['Müşteri memnuniyeti', 'Yıllık NPS = 2 × yılın ortalama hizmet skoru − 100.', 'Ortalama hizmet skoru'])
section(ws, 'Sahne göstergeleri', 3)
for a, b in [('Brüt prim', 'Kümülatif yazılan prim; sahne bu değere göre sıralanır.'), ('Pazar payı', 'Brüt prim payı.'),
             ('Kâr / zarar', 'Kümülatif teknik kâr.'), ('Hasar / prim', 'Kümülatif brüt hasar ÷ brüt prim.'), ('NPS', 'O ana kadarki ortalama hizmet skorundan.')]:
    row(ws, [a, b, ''])
section(ws, 'Arka plandaki birleşik puan (yalnızca hesap dosyasında)', 3)
labels = ['Kârlılık', 'Pazar payı', 'Memnuniyet']
for i, w in enumerate(D['weights']):
    row(ws, [f'Ağırlık · {labels[i]}', f'%{w}', ''])
row(ws, ['Kârlılık eşikleri', f"Kâr/sermaye %{A['profitFloor']['value'] * 100:.0f} → 0 puan, %{A['profitTarget']['value'] * 100:.0f} → 100 puan (sezon içinde geçen aya orantılı).", ''])
row(ws, ['Pay eşiği', f"Adil payın (1/takım) {tr(A['shareTarget']['value'], 1)} katı → 100 puan.", ''])
row(ws, ['Memnuniyet eşiği', f"Ortalama hizmet skoru {A['serviceTarget']['value']} → 100 puan.", ''])

# ——— 12. Hesap adımları ———
ws = sheet('Hesap adımları', ['Adım', 'Ne olur', 'Formül'], [8, 50, 90], 'Motor her ay, her hücre ve her takım için bu adımları sırayla çalıştırır.')
steps = [
    ('Müşteri havuzu', 'Hücrenin o ay pazara çıkan müşterisi.', 'yıllık poliçe × hücre payı ÷ 12 × mevsimsellik × olay talep çarpanı × (0,985 + şans × 0,03)'),
    ('Teklif', 'Takımın o hücreye fiyatı.', 'baz prim × beş segment katsayısının çarpımı'),
    ('Seçim ağırlığı', 'Müşterinin takımı tercih etme gücü.', '(teklif ÷ referans prim)^(−β) × görünürlük^sadakat × e^(hizmet ağırlığı × itibar × (hizmet − 98) ÷ 40); görünürlük pazarlama × (1 − kampanya %) ile hesaplanır'),
    ('Satış', 'Havuzun takımlar ve dış seçenek arasında paylaşılması.', 'havuz × ağırlık ÷ (dış seçenek + Σ ağırlık)'),
    ('Kampanya erişimi', 'Ortak hedef kitlede gösterim ve erişim.', 'gösterim = medya ÷ CPM × 1.000; erişim = gösterim × min(1, kitle ÷ Σ gösterim); frekans = max(1, Σ gösterim ÷ kitle)'),
    ('Kampanya müşterisi', 'Huniden geçen ve hediye bütçesine sığan müşteri.', 'erişim × ilgi × tıklama × hit (frekans > 5 ise ×1,1) × dijital hücre payı × mevsimsellik × olay talebi × min(2, (teklif ÷ referans)^(−β)); en fazla hediye bütçesi ÷ hediye maliyeti'),
    ('Prim ve kanal gideri', 'Yazılan prim ve satış maliyeti.', 'satış × teklif; kanal gideri = prim × kanal gider oranı'),
    ('Beklenen hasar', 'Satılan poliçelerin beklenen adedi ve maliyeti.', 'satış × frekans × olay hasar çarpanı; maliyet = adet × şiddet'),
    ('Kapasite ve hizmet', 'Gelen dosya ile kapasitenin karşılaştırılması.', 'kapasite = hasar op. ÷ 12 ÷ (0,1 × şiddet); hizmet = 98 − 60 × max(0, kullanım − 0,8), 30–98 arası'),
    ('Hasar çekimi', 'Gerçekleşen hasar.', 'adet ~ Poisson(normal yaklaşımı), tutar ~ Gamma; sızıntı = 1 + 0,15 × (kullanım − 1)'),
    ('Reasürans', 'Devredilen prim ve hasar, komisyon.', 'devir = prim × %30; komisyon = devir × %25'),
    ('Kâr', 'Ayın sonunda kümülatif teknik kâr.', 'net prim − net hasar − (kanal gideri + sabit gider ÷ 12 + karar bütçesi ÷ 12 − komisyon)'),
]
for i, (a, b, c) in enumerate(steps, 1):
    row(ws, [i, f'{a}: {b}', c], [N0])

# ——— 13. Pazar hücreleri ———
ws = sheet('Pazar hücreleri', [DN['city'], DN['channel'], DN['vehicle'], DN['persona'], DN['type'], 'Örneklem adedi', 'Yıllık poliçe', 'Son dönem ort. prim',
                               'Rakip fiyat endeksi', 'Frekans', 'Ort. hasar tutarı', 'Beklenen hasar / poliçe', 'Referans prim', 'Beklenen hasar/prim'],
           [12, 10, 14, 14, 12, 11, 11, 13, 11, 9, 12, 12, 12, 12],
           'Her satır bir segment hücresi. Referans prim, piyasanın o hücreye uyguladığı fiyattır (takımların teklifi bununla kıyaslanır).')
for x in cells:
    c = x['cell']
    row(ws, [LV['city'][c[0]]['name'], LV['channel'][c[1]]['name'], LV['vehicle'][c[2]]['name'], LV['persona'][c[3]]['name'], LV['type'][c[4]]['name'],
             c[5], c[5] / D['customers'] * D['policies'], c[6], c[7], x['freq'], x['sev'], x['cost'], x['reference'], x['cost'] / x['reference']],
        [None, None, None, None, None, N0, N0, EUR2, N4, PCT1, EUR, EUR2, EUR2, PCT1])
ws.auto_filter.ref = f"A{ws.freeze_panes[1:] and int(ws.freeze_panes[1:]) - 1}:N{ws._next - 1}"

# ——— Veri kullanımı ———
ws = sheet('Veri kullanımı', ['Sütun (pricing_case_data · DATA)', 'Oyunda', 'Nasıl'], [32, 16, 100],
           'Vaka verisinin 26 sütunu ve oyunun onları nasıl kullandığı. Verinin bir örneklem olduğunu unutmayın: oyun profili kullanır, satırları tek tek oynatmaz.')
for col, used, how in [
    ('customer_id', 'Hayır', 'Yalnızca satır kimliği.'),
    ('renewal_channel', 'Evet', 'Kanal boyutu (hücre). Not: oyun yenileme değil, sıfırdan yeni iş yarışıdır.'),
    ('city', 'Evet', 'İl boyutu (hücre).'), ('vehicle_segment', 'Evet', 'Araç yaşı boyutu (hücre).'),
    ('customer_type', 'Evet', 'Müşteri tipi boyutu (hücre).'), ('persona', 'Evet', 'Persona boyutu (hücre) ve davranış (fiyat/hizmet/kanal duyarlılığı).'),
    ('base_premium_prev_term', 'Bilgi', 'Hücre ortalaması olarak saklanır (hesap dosyasında görünür); hesaba girmez.'),
    ('offer_premium_gross', 'Hayır', 'Yenileme teklifi; oyunda yenileme yok.'), ('price_change_pct', 'Hayır', 'Yenileme fiyat değişimi; oyunda karşılığı yok.'),
    ('discount_offered', 'Hayır', 'İndirim kararı; oyunda indirim kaldıracı yok (fiyat katsayılarla verilir).'), ('discount_pct', 'Hayır', 'Aynı.'),
    ('final_premium_gross', 'Evet', 'Kalibrasyon: piyasa referans primi bu ortalamayı (€294,67) yeniden üretir.'),
    ('competitor_price_index', 'Evet', 'Hücrenin referans primine çarpan olarak girer.'),
    ('claims_count_12m', 'Hayır', 'Hasar modelden üretilir. Bu sütun eski şiddetle (€133) üretildiği için oyunla tutarsız.'),
    ('claims_paid_12m', 'Hayır', 'Aynı; gerçekleşen hasar/prim ≈ %4,5 gösterir, oyun ≈ %60.'),
    ('freq_coef_combined', 'Kalibrasyon', 'Frekans normalizasyon sabiti buradan ölçülür. Takımlara giderse cevabı verir.'),
    ('sev_coef_combined', 'Kalibrasyon', 'Şiddet normalizasyon sabiti buradan ölçülür. Takımlara giderse cevabı verir.'),
    ('premium_coef_combined', 'Hayır', 'Piyasa prim katsayıları COEFFICIENTS sayfasından okunur.'),
    ('poisson_lambda … expected_combined_ratio', 'Hayır', 'Modelin türetilmiş çıktıları; takımlara giderse fiyatlamayı doğrudan çözer.'),
]:
    row(ws, [col, used, how])

# ——— Denetim bulguları ———
ws = sheet('Denetim bulguları', ['#', 'Konu', 'Bulgu', 'Etki', 'Durum', 'Öneri'], [5, 26, 70, 12, 18, 60],
           'Veri, model konfigürasyonu, girdi dosyaları ve oyun mantığı karşılaştırıldı. “Düzeltildi” oyunda çözüldü; “Sizin dosyanız” veri/girdi dosyasında düzeltilmeli; “Karar” sizin tasarım kararınız.')
findings = [
    ('Hasar tutarı', 'Model konfigürasyonunda temel şiddet €133; verideki primler €295 olduğundan gerçekleşen hasar/prim %4,5 çıkıyor.', 'Yüksek', 'Oyunda düzeltildi', 'Konfigürasyonda Base Severity = 1753,90 yapın ve verideki hasar sütunlarını yeniden üretin.'),
    ('Verideki hasar sütunları', 'claims_count_12m / claims_paid_12m eski şiddetle üretilmiş; takımlar analizde %4,5 hasar/prim görür, oyun %60 oynar.', 'Bilgi', 'Düzeltildi', 'veri/pricing_case_data_duzeltilmis.xlsx (hasar tutarları ×13,16; gerçekleşen hasar/prim %59,4). Katılımcıya veri/katilimci_verisi.xlsx verilir.'),
    ('Cevabı veren sütunlar', 'P–Z sütunları (katsayı çarpımları, beklenen hasar/prim) fiyatlamayı doğrudan çözer.', 'Bilgi', 'Düzeltildi', 'Katılımcı dosyasında yalnızca A–O sütunları ve bir sözlük sayfası var.'),
    ('İl prim katsayıları', 'Konfigürasyondaki CITY prim katsayıları verideki COEFFICIENTS ile farklı (ör. Ankara 0,945 / 1,018).', 'Orta', 'Oyunda düzeltildi', 'Konfigürasyonu verideki değerlerle eşitleyin.'),
    ('Yenileme verisi, yeni iş oyunu', 'Veri bir yenileme portföyü (önceki prim, fiyat değişimi, indirim); oyun herkesin sıfırdan yarıştığı yeni iş pazarı. İndirim/yenileme kaldıracı yok.', 'Orta', 'Karar', 'Brifingde açıkça söyleyin: indirim ve fiyat değişimi sütunları bağlam içindir.'),
    ('Banka müşterisi personası', 'Banka müşterilerinin yalnızca %20’si banka kanalından, %61’i acenteden geliyor; persona en yüksek kanal sadakatine (1,4) sahip.', 'Düşük', 'Karar', 'Sadakat “her kanalda görünürlüğe hassas” anlamında çalışıyor; isim yanıltıcıysa davranışı ya da adı gözden geçirin.'),
    ('Kampanya hedef kitlesi', '“18–55 yaş dijital, Joyful Disregarders” veride yok (yaş ve bu persona yok). Kampanya müşterileri dijital kanal hücrelerine örneklem payıyla dağıtılıyor.', 'Orta', 'Varsayım', 'Hedef kitleyi bir personaya bağlamak isterseniz söyleyin.'),
    ('Kampanya hacmi', f"Örnek stratejilerle kampanya yılda ~14 bin müşteri getiriyor; bu organik dijital pazarın (~{D['digitalPolicies']:,.0f}) ~%56’sı.".replace(',', '.'), 'Orta', 'Varsayım', 'Fazla bulursanız Kural stüdyosunda hedef kitle payını ya da hediye hit oranlarını düşürün.'),
    ('Medya maliyeti', 'Dosyada “1 kullanıcıya gösterme maliyeti €10” yazıyor; oyunda bin gösterim başına €10 (CPM) kullanılıyor.', 'Orta', 'Karar (CPM)', 'Dosyadaki metni “1.000 gösterim başına” olarak düzeltin.'),
    ('Kampanya ve olaylar', 'Dijital rakip / sıfır araç kredisi olayları kampanya kitlesini etkilemiyordu; mevsimsellik de uygulanmıyordu.', 'Orta', 'Oyunda düzeltildi', '—'),
    ('Katsayı aralığı', 'Oyun 0,50–2,50 (0,05 adım) kullanıyor; verideki risk katsayıları (ortalamaya göre dijital ~0,70 … broker ~1,75, ticari ~1,50) bu aralığa sığıyor. Marketing_Input dosyasındaki 0,80–1,20 notu artık geçerli değil.', 'Bilgi', 'Oyunda ayarlandı', 'Elinizdeki eski dosyada aralığı güncelleyin ya da sitedeki şablonu dağıtın.'),
    ('Karar dosyası', 'Oyunun indirdiği şablon artık Marketing_Input tasarımında (Input, Premium, Marketing, Claim). Claim sayfası oyunun yapısında: hasar operasyonu (€) ve kota paylı var/yok; Marketing sayfasına pazarlama bütçesi, kampanya payı, kanal tablosu ve seçilen hediye eklendi.', 'Bilgi', 'Oyunda yapıldı', 'Takımlara sitedeki şablonu dağıtın; elinizdeki eski input dosyaları oyuna uymaz.'),
    ('Girdi dosyası: metinler', '“Interest Rate” aslında ilgi oranı; F14’teki “ad only once” notu yarım. Oyunun şablonunda bunlar dosyadaki gibi bırakıldı, yalnızca medya maliyeti “1.000 gösterim” olarak düzeltildi.', 'Düşük', 'Sizin dosyanız', 'İsterseniz metinleri düzeltin, şablona da yansıtırım.'),
    ('Sermaye açıklaması', 'Açıklama “sıfırın altına düşen şampiyon olamaz” diyordu; özkaynak kuralı kapalı.', 'Düşük', 'Oyunda düzeltildi', '—'),
    ('Kârlılık dengesi', 'Örnek botlar ucuz fiyatladığı için zarar ediyor; aynı sahaya karşı pahalı fiyat (hedef hasar/prim %45) + %40 kampanya ile €1,19 mn kâr mümkün.', 'Bilgi', 'Kontrol edildi', 'Kâr kupası anlamlı; sabit gider ve bütçe adil dilimin %10’u.'),
    ('Birleşik puan', 'Sahne ve final üç kupa kullanıyor; ağırlıklı puan yalnızca hesap dosyasında duruyor.', 'Bilgi', 'Bilerek', 'İsterseniz tamamen kaldırılabilir.'),
]
for i, f in enumerate(findings, 1):
    row(ws, [i, *f], [N0])

wb.save('oyun-varsayimlari.xlsx')
print(f'oyun-varsayimlari.xlsx · {len(wb.sheetnames)} sayfa · {len(cells)} hücre')
