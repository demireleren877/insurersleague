"""Rebuilds the case data with claim amounts that match its premiums, and the participant version.

The original pricing_case_data.xlsx drew claim amounts with a base severity of €133.244 while its premiums
imply ~€1,754 (portfolio loss ratio ~60%, the data's own expected LR). Claim counts were right; only the
amounts were ~13× too small (observed loss ratio 4.5%).

Fix: every claim amount is multiplied by 1,753.90 / 133.244. Scaling a Gamma variable scales its scale
parameter, so this is exactly the same draw made with the corrected severity; zero-claim rows stay zero.

Outputs (in veri/):
  pricing_case_data_duzeltilmis.xlsx   all sheets and columns, corrected (moderator copy)
  katilimci_verisi.xlsx                what teams get: only the columns the game uses (segments, premium,
                                       competitor index, corrected claims), a column dictionary and the
                                       marketing inputs — no model columns, no coefficient tables

Usage:  python3 scripts/fix-case-data.py [path/to/pricing_case_data.xlsx]
"""
import json
import subprocess
import sys
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else 'veri/pricing_case_data.xlsx'
OLD_SEV, NEW_SEV = 133.244, 1753.90
K = NEW_SEV / OLD_SEV

HEAD_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='FF1F4E78')
EUR2, EUR0, EUR3, PCT1 = '\\€#,##0.00', '\\€#,##0', '\\€#,##0.000', '0.0%'

src = load_workbook(SRC, read_only=True, data_only=False)
rows = list(src['DATA'].iter_rows(values_only=True))
head, body = list(rows[0]), [list(r) for r in rows[1:] if r and r[0]]
col = {h: i for i, h in enumerate(head)}

# ——— Corrected claims ———
for r in body:
    r[col['claims_paid_12m']] = round(float(r[col['claims_paid_12m']]) * K, 2)
    r[col['expected_severity_eur']] = round(float(r[col['expected_severity_eur']]) * K, 2)
    r[col['gamma_scale']] = round(float(r[col['gamma_scale']]) * K, 4)

DATA_FMT = {'base_premium_prev_term': EUR2, 'offer_premium_gross': EUR2, 'price_change_pct': PCT1, 'discount_pct': PCT1, 'final_premium_gross': EUR2,
            'claims_paid_12m': EUR2, 'p_1plus_claim': PCT1, 'expected_severity_eur': EUR2, 'gamma_scale': EUR2, 'expected_loss_ratio': PCT1,
            'expense_ratio': PCT1, 'expected_combined_ratio': PCT1}
DATA_W = {'customer_id': 13, 'renewal_channel': 17, 'city': 12, 'vehicle_segment': 17, 'customer_type': 15, 'base_premium_prev_term': 24, 'offer_premium_gross': 21,
          'price_change_pct': 18, 'discount_offered': 16, 'discount_pct': 14, 'final_premium_gross': 21, 'competitor_price_index': 24, 'claims_count_12m': 18,
          'claims_paid_12m': 17, 'persona': 14, 'freq_coef_combined': 20, 'sev_coef_combined': 19, 'premium_coef_combined': 23, 'poisson_lambda': 16,
          'p_1plus_claim': 15, 'expected_severity_eur': 23, 'gamma_shape': 13, 'gamma_scale': 13, 'expected_loss_ratio': 21, 'expense_ratio': 15, 'expected_combined_ratio': 25}


def write_table(ws, header, data, fmts=None, widths=None, filt=True):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical='center')
    for r in data:
        ws.append(r)
    for i, h in enumerate(header, 1):
        letter = get_column_letter(i)
        if widths and h in widths:
            ws.column_dimensions[letter].width = widths[h]
        if fmts and h in fmts:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmts[h]
    ws.freeze_panes = 'A2'
    if filt:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(data) + 1}'


# ——— 1. Full corrected workbook ———
full = Workbook()
ws = full.active
ws.title = 'DATA'
write_table(ws, head, body, DATA_FMT, DATA_W)

coef_rows = [list(r) for r in src['COEFFICIENTS'].iter_rows(values_only=True)]
ci = {h: i for i, h in enumerate(coef_rows[0]) if h}
for r in coef_rows[1:]:
    if r[ci['Target Severity EUR']] is not None:
        r[ci['Target Severity EUR']] = round(float(r[ci['Target Severity EUR']]) * K, 3)
ws = full.create_sheet('COEFFICIENTS')
write_table(ws, coef_rows[0], coef_rows[1:], {'Target Share': PCT1, 'Target Freq': PCT1, 'Target Severity EUR': EUR3, 'Expense Ratio': PCT1},
            {'Variable': 12, 'Level': 14, 'Target Share': 14, 'Target Freq': 13, 'Freq Coef': 12, 'Target Severity EUR': 21, 'Severity Coef': 15, 'Premium Coef': 14, 'Expense Ratio': 15, 'Source': 34}, filt=False)

# VALIDATION, recomputed from the corrected rows (same layout and order as the original sheet)
orig_val = [r for r in src['VALIDATION'].iter_rows(values_only=True)]
dims = {'City': 'city', 'Channel': 'renewal_channel', 'Vehicle Segment': 'vehicle_segment', 'Persona': 'persona', 'Customer Type': 'customer_type'}
agg = defaultdict(lambda: [0, 0, 0.0, 0.0, 0.0])
for r in body:
    keys = [(v, r[col[c]]) for v, c in dims.items()] + [('TOTAL', 'TOTAL')]
    for key in keys:
        a = agg[key]
        a[0] += 1; a[1] += r[col['claims_count_12m']]; a[2] += r[col['claims_paid_12m']]; a[3] += r[col['final_premium_gross']]; a[4] += r[col['expected_combined_ratio']]
n_all = len(body)
val = []
for r in orig_val[1:]:
    if not r or r[0] is None:
        continue
    n, cnt, paid, prem, cr = agg[(r[0], r[1])]
    i = len(val) + 2
    val.append([r[0], r[1], n, n / n_all, cnt, cnt / n, round(paid, 2), paid / cnt if cnt else 0, round(prem, 2), paid / prem, cr / n, f'=I{i}/C{i}'])
ws = full.create_sheet('VALIDATION')
write_table(ws, ['Variable', 'Level', 'Policies', 'Share', 'Claim Count', 'Observed Frequency', 'Claim Paid EUR', 'Observed Severity EUR', 'Premium EUR', 'Observed LR', 'Avg Expected CR', 'Avg Premium EUR'], val,
            {'Share': PCT1, 'Observed Frequency': PCT1, 'Claim Paid EUR': EUR2, 'Observed Severity EUR': EUR2, 'Premium EUR': EUR0, 'Observed LR': PCT1, 'Avg Expected CR': PCT1, 'Avg Premium EUR': EUR2},
            {'Variable': 14, 'Level': 14, 'Policies': 10, 'Share': 9, 'Claim Count': 13, 'Observed Frequency': 20, 'Claim Paid EUR': 16, 'Observed Severity EUR': 23, 'Premium EUR': 14, 'Observed LR': 13, 'Avg Expected CR': 17, 'Avg Premium EUR': 17}, filt=False)

logic = [list(r) for r in src['MODEL_LOGIC'].iter_rows(values_only=True) if r and r[0]]
logic.append(['Base Severity EUR', f'{NEW_SEV:,.2f} — calibrated to the premium level (the earlier value {OLD_SEV} gave a 4.5% loss ratio against these premiums). Claim amounts = earlier amounts × {K:.6f}, i.e. the same Gamma draws at the corrected scale.'])
ws = full.create_sheet('MODEL_LOGIC')
write_table(ws, logic[0], logic[1:], widths={'Output': 30, 'Formula / Method': 120}, filt=False)
full.save('veri/pricing_case_data_duzeltilmis.xlsx')

# ——— 2. Participant workbook ———
# Only what the game uses: the five segment columns, the premium the customer pays, the competitors'
# price index (it sets the rest of the market's price) and the claims. The renewal-offer columns
# (previous premium, offer, price change, discount) play no part in the game.
keep = ['customer_id', 'renewal_channel', 'city', 'vehicle_segment', 'customer_type', 'final_premium_gross', 'competitor_price_index', 'claims_count_12m', 'claims_paid_12m', 'persona']
part = Workbook()
ws = part.active
ws.title = 'DATA'
write_table(ws, keep, [[r[col[k]] for k in keep] for r in body], DATA_FMT, DATA_W)
ws = part.create_sheet('DICTIONARY')
dictionary = [
    ('customer_id', 'Customer identifier', 'Müşteri kimliği'),
    ('renewal_channel', 'Sales channel: Agency, Bank, Digital, Broker', 'Satış kanalı: acente, banka, dijital, broker'),
    ('city', 'City: Istanbul, Ankara, Izmir, Bursa, Antalya, Other', 'İl: İstanbul, Ankara, İzmir, Bursa, Antalya, diğer'),
    ('vehicle_segment', 'Vehicle age: New/0-1y, Mid (2-6y), Old (7y+)', 'Araç yaşı: yeni (0–1), orta (2–6), eski (7+)'),
    ('customer_type', 'Individual or Commercial', 'Bireysel ya da ticari'),
    ('final_premium_gross', 'Gross premium the customer pays (EUR)', 'Müşterinin ödediği brüt prim (EUR)'),
    ('competitor_price_index', 'Competitors’ price ÷ our price (1.05 = competitors 5% dearer)', 'Rakip fiyatı ÷ bizim fiyat (1,05 = rakipler %5 pahalı)'),
    ('claims_count_12m', 'Number of claims in the last 12 months', 'Son 12 aydaki hasar adedi'),
    ('claims_paid_12m', 'Claims paid in the last 12 months (EUR)', 'Son 12 ayda ödenen hasar (EUR)'),
    ('persona', 'Customer persona: Bank customer, Post-claim, Price-driven, Value-driven', 'Müşteri personası: banka müşterisi, hasar sonrası, fiyat odaklı, değer odaklı'),
]
ws.append(['Column', 'Description', 'Açıklama'])
for c in ws[1]:
    c.font, c.fill = HEAD_FONT, HEAD_FILL
for r in dictionary:
    ws.append(list(r))
for letter, w in zip('ABC', (24, 70, 70)):
    ws.column_dimensions[letter].width = w
ws.freeze_panes = 'A2'
ws.append([])
ws.append(['Note', f'The data is a sample of {n_all:,} customers from the casco market; its profile matches the market the game simulates.', f'Veri kasko pazarından {n_all:,} müşterilik bir örneklemdir; profili oyunun simüle ettiği pazarla aynıdır.'.replace(',', '.')])

# The marketing inputs (Marketing_Input.xlsx) with this game's own budget and audience.
game = json.loads(subprocess.run(['node', '--input-type=module', '-e', """
import { scenario, cascoMoney, campaignRules, campaignAudience, offerName } from './dist/engine.js';
const c = scenario('tr'), m = cascoMoney(c), R = c.rules, K = campaignRules(R);
console.log(JSON.stringify({ budget: m.budget, audience: campaignAudience(R, m), K, names: K.offers.map(o => [offerName(o.id, 'en'), offerName(o.id, 'tr')]) }));
"""], capture_output=True, text=True, check=True).stdout)
KC = game['K']
ws = part.create_sheet('MARKETING')
ws.append(['Item', 'Value', 'Açıklama / Description'])
for c in ws[1]:
    c.font, c.fill = HEAD_FONT, HEAD_FILL
rows = [
    ('Population', 84_000_000, 'Nüfus'),
    ('18–55 digital users', KC['digitalUsers'], '18–55 yaş dijital kullanıcı'),
    ('Target group: Joyful Disregarders', KC['targetShare'], 'Hedef kitle: dijital kullanıcıların payı'),
    ('Marketing budget per team (EUR, at most)', game['budget'], 'Takım başına pazarlama bütçesi (en fazla). Medya ve hediye arasında bölünür.'),
    ('Target group per team in a year', round(game['audience']), 'Bu oyunda takım başına yıllık hedef kitle (vakadaki euro başına kişi korunur)'),
    ('Media cost per 1,000 impressions (EUR)', KC['cpm'], '1.000 reklam gösteriminin maliyeti'),
    (f"Frequency bonus: seen more than {KC['frequency']} times", KC['frequencyBonus'], f"Reklamı ortalama {KC['frequency']} kereden fazla gören kitlede hit oranı bu kadar artar"),
    ('Price effect cap on conversion (×)', KC['priceCap'], 'Ucuz fiyat kampanya dönüşümünü en fazla bu kat artırır'),
]
for r in rows:
    ws.append(list(r))
for i, fmt in enumerate(['#,##0', '#,##0', PCT1, EUR0, '#,##0', EUR2, PCT1, '0.0'], start=2):
    ws.cell(row=i, column=2).number_format = fmt
ws.append([])
ws.append(['Gift', 'Interest rate', 'Cost (EUR)', 'Click rate', 'Hit ratio', 'Hediye'])
for c in ws[ws.max_row]:
    c.font, c.fill = HEAD_FONT, HEAD_FILL
for o, (en, tr) in zip(KC['offers'], game['names']):
    ws.append([en, o['interest'], o['cost'], o['click'], o['hit'], tr])
    r = ws.max_row
    for cl, fmt in (('B', PCT1), ('C', EUR0), ('D', PCT1), ('E', PCT1)):
        ws[f'{cl}{r}'].number_format = fmt
ws.append([])
ws.append(['How it works', '', 'Customers per gift = reach × gift weight × interest × click × hit ratio; each gift serves at most (gift budget × its weight) ÷ its cost customers. Every acquired customer gets one gift.'])
ws.append(['Nasıl çalışır', '', 'Hediye başına müşteri = erişim × hediye ağırlığı × ilgi × tıklama × hit oranı; her hediye en fazla (hediye bütçesi × ağırlığı) ÷ maliyeti kadar müşteriye yeter. Kazanılan her müşteri bir hediye alır.'])
for letter, w in zip('ABCDEF', (40, 16, 90, 12, 12, 30)):
    ws.column_dimensions[letter].width = w
ws.freeze_panes = 'A2'
part.save('veri/katilimci_verisi.xlsx')

paid = sum(r[col['claims_paid_12m']] for r in body)
prem = sum(r[col['final_premium_gross']] for r in body)
print(f'{n_all} rows · claims ×{K:.4f} · observed loss ratio {paid / prem:.1%} · participant columns: {len(keep)}')
