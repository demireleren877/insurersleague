"""Builds simulasyon-ornegi.xlsx: one simulated season, laid out so the logic can be questioned line by line,
and simulasyon-hesap.xlsx: the same season as the formula workbook (every number recomputed by Excel).

Six sample teams play the default game; one of them changes its plan at the second quarter review, so the
quarter mechanics show too. Run again after rule or data changes:  python3 scripts/build-simulation-book.py
"""
import json, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DUMP = r"""
import { writeFileSync } from 'node:fs';
const e = await import('./dist/engine.js');
const c = await import('./dist/casco.js');
const g = await import('./dist/js/game.js');
const { buildAuditWorkbook } = await import('./dist/js/audit.js');
const cfg = e.scenario('tr'), R = e.rulesOf(cfg), money = e.cascoMoney(cfg);
const teams = e.teams('tr', cfg);
// One quarter change: the actuary team starts a restaurant-card campaign from month 7.
const changed = teams[0];
changed.strategy = { ...changed.strategy, campaign: 0 };
changed.strategyHistory = [
  { effectiveMonth: 0, strategy: structuredClone(changed.strategy) },
  { effectiveMonth: 6, strategy: { ...structuredClone(changed.strategy), campaign: 40, mediaShare: 25, offer: 'restaurant' } }
];
const months = c.simulateCasco(teams, cfg, { trace: true });
const state = g.freshSession(0, { code: '000000', lang: 'tr' });
state.teams = teams;
writeFileSync('simulasyon-hesap.xlsx', await buildAuditWorkbook(state, { lang: 'tr' }));
const dims = e.DIMENSIONS;
console.log(JSON.stringify({
  teams: teams.map(t => ({ id: t.id, name: t.name, strategy: t.strategy, history: t.strategyHistory ?? null })),
  months: months.map(m => ({ ...m, rows: m.rows })),
  cells: e.marketCells().map(cell => ({ cell, ...e.cellRisk(cell, R) })),
  levels: Object.fromEntries(dims.map(d => [d, R.dimensions[d].map((_, i) => e.levelName(d, i, 'tr'))])),
  dimNames: Object.fromEntries(dims.map(d => [d, e.dimensionName(d, 'tr')])),
  behavior: R.behavior, commercialPrice: R.commercialPrice,
  money, policies: e.policiesOf(cfg), customers: 50000, seed: cfg.seed, weights: cfg.weights,
  offers: Object.fromEntries(e.campaignRules(R).offers.map(o => [o.id, { ...o, name: e.offerName(o.id, 'tr') }])),
  campaign: e.campaignRules(R), reinsurance: R.reinsurance, service: R.service, expense: R.dimensions.channel.map(x => x.expense),
  months_tr: e.monthsOf('tr')
}));
"""
D = json.loads(subprocess.run(['node', '--input-type=module', '-e', DUMP], capture_output=True, text=True, check=True).stdout)
T, MO, CELLS, LV, DN, MONEY = D['teams'], D['months'], D['cells'], D['levels'], D['dimNames'], D['money']
NAME = {t['id']: t['name'] for t in T}
MN = D['months_tr']

INK, MUTED, LINE = '1E1E1E', '5B6470', 'C3C8CE'
HEAD, SECTION, NOTE = PatternFill('solid', fgColor='1E1E1E'), PatternFill('solid', fgColor='FFE066'), PatternFill('solid', fgColor='FFF6D6')
GOOD, BAD = PatternFill('solid', fgColor='D3F9D8'), PatternFill('solid', fgColor='FFE3E3')
thin = Border(bottom=Side(style='thin', color=LINE))
EUR, EUR2, PCT, PCT1, N0, N1, N2, N4 = '#,##0 "€"', '#,##0.00 "€"', '0%', '0.0%', '#,##0', '#,##0.0', '0.00', '0.0000'

wb = Workbook()
wb.remove(wb.active)


def sheet(title, headers, widths, intro=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, title).font = Font(bold=True, size=16, color=INK)
    r = 2
    if intro:
        ws.cell(r, 1, intro).font = Font(size=11, color=MUTED, italic=True)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, len(headers)))
        ws.row_dimensions[r].height = max(30, 15 * (len(intro) // 140 + 1))
        r += 1
    r += 1
    for i, h in enumerate(headers, 1):
        c = ws.cell(r, i, h)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 32
    ws.freeze_panes = ws.cell(r + 1, 1)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws._head = r
    ws._next = r + 1
    return ws


def row(ws, values, fmts=None, fill=None, bold=False):
    r = ws._next
    for i, v in enumerate(values, 1):
        c = ws.cell(r, i, v)
        c.alignment = Alignment(wrap_text=isinstance(v, str) and len(v) > 30, vertical='top')
        c.border = thin
        if fmts and i - 1 < len(fmts) and fmts[i - 1]:
            c.number_format = fmts[i - 1]
        if fill:
            c.fill = fill
        if bold:
            c.font = Font(bold=True)
    ws._next = r + 1


def filt(ws, last_col):
    ws.auto_filter.ref = f'A{ws._head}:{get_column_letter(last_col)}{ws._next - 1}'


row_of = lambda m, tid: next(r for r in MO[m]['rows'] if r['id'] == tid)
final = {r['id']: r for r in MO[11]['rows']}
plan_at = lambda t, m: next((h['strategy'] for h in sorted(t['history'] or [], key=lambda h: -h['effectiveMonth']) if h['effectiveMonth'] <= m), t['strategy'])

# ——— Oku beni ———
ws = sheet('Oku beni', ['Konu', 'Açıklama'], [28, 120],
           'Varsayılan kurallarla oynatılmış tek bir sezon. Amaç: motorun her adımını rakamlarla sorgulayabilmek. Aynı sezonun formül formül hesabı simulasyon-hesap.xlsx dosyasındadır (Excel her sayıyı yeniden hesaplar, Puan sayfasında fark 0,00 olmalı).')
for a, b in [
    ('Takımlar', '6 örnek strateji: aktüer, düz tarife, hacim, marj (reasüranslı), dijital (kampanyalı), hizmet. Hepsi aynı veriyi farklı okuyor.'),
    ('Çeyrek değişikliği', f"{NAME[T[0]['id']]} kampanyasız başlıyor, 7. aydan itibaren restoran kartıyla %40 kampanya yürütüyor (2. çeyrek molası kararı)."),
    ('Pazar', f"Yıllık {D['policies']:,} poliçe, 50.000 kişilik örneklemin profiliyle. Rastgelelik tohumu {D['seed']}.".replace(',', '.')),
    ('Muhasebe', 'Yazım yılı esası: bir ayda satılan poliçe tam yıllık primini ve nihai hasarını o ay kaydeder. Bu yüzden aylık rakamlar “o ay yazılan iş”tir.'),
    ('Takım planları', 'Her takımın kararları; çeyrek değişikliği ayrı sütunda.'),
    ('Aylık sonuçlar', 'Takım × ay: satış, prim, hasar, oranlar, kapasite, hizmet, NPS, pay, kâr.'),
    ('Gider dökümü', 'Yıl sonu kâr köprüsü: prim → devir → hasar → giderler → kâr. Son sütun motorun kârıyla farkı gösterir (0 olmalı).'),
    ('Kampanya hunisi', 'Takım × ay: medya, gösterim, erişim, frekans, aday, istenen ve kazanılan müşteri, hediye kullanımı ve israfı.'),
    ('Pazar (aylık)', 'Ay başına pazar havuzu, takımların sattığı, kimseden almayanlar, olaylar ve endeksler.'),
    ('Segment sonuçları', 'Takım × hücre yıllık: pazar ve kampanya poliçesi, teklif ÷ referans, prim, beklenen hasar, beklenen hasar/prim. Filtrelenebilir.'),
    ('Final', 'Brüt prim sıralaması ve üç kupa.'),
    ('Mantık kontrolleri', 'Motorun tutarlılığını sınayan otomatik kontroller (kâr kimliği, payların toplamı, hediye sınırı, hizmet formülü…).'),
]:
    row(ws, [a, b])

# ——— Takım planları ———
heads = ['Karar'] + [t['name'] for t in T] + [f"{NAME[T[0]['id']]} (7. aydan)"]
ws = sheet('Takım planları', heads, [34] + [18] * (len(T) + 1))
def plan_rows(st):
    return [
        ('Baz prim', st['basePremium'], EUR2),
        *[(f"{DN[d]} · {LV[d][i]}", st['coef'][d][i], N2) for d in ['city', 'channel', 'vehicle', 'persona', 'type'] for i in range(len(LV[d]))],
        ('Pazarlama bütçesi', st['marketing'], EUR),
        *[(f"Pazarlama odağı · {LV['channel'][i]}", st['channelFocus'][i] / 100, PCT) for i in range(4)],
        ('Kampanya payı (pazarlamanın)', (st.get('campaign') or 0) / 100, PCT),
        ('Medya payı (kampanyanın)', (st.get('mediaShare') or 0) / 100, PCT),
        ('Kampanya hediyesi', D['offers'][st.get('offer') or 'concert']['name'], None),
        ('Hasar operasyonu', st['claimsOps'], EUR),
        ('Reasürans', 'Var' if st['reinsurance'] else 'Yok', None),
        ('Kullanılan bütçe', st['marketing'] + st['claimsOps'] + (MONEY['reinsuranceFee'] if st['reinsurance'] else 0), EUR),
    ]
labels = [x[0] for x in plan_rows(T[0]['strategy'])]
cols = [plan_rows(t['strategy']) for t in T] + [plan_rows(T[0]['history'][1]['strategy'])]
for k, label in enumerate(labels):
    row(ws, [label] + [c[k][1] for c in cols], [None] + [c[k][2] for c in cols])

# ——— Aylık sonuçlar ———
H = ['Takım', 'Ay', 'Yeni poliçe', '…kampanyadan', 'Kümülatif poliçe', 'Ayın brüt primi', 'Kümülatif brüt prim', 'Ayın hasarı', 'Kümülatif hasar adedi',
     'Ayın hasar/primi', 'Kümülatif hasar/prim (brüt)', 'Net hasar/prim', 'Bileşik oran', 'Kapasite (dosya/ay)', 'Gelen dosya/ay', 'Kullanım',
     'Ayın hizmet skoru', 'Ort. hizmet', 'NPS', 'Pazar payı (prim)', 'Ayın prim payı', 'Poliçe payı', 'Ayın kârı', 'Kümülatif kâr', 'Özkaynak']
ws = sheet('Aylık sonuçlar', H, [16, 6] + [13] * (len(H) - 2), 'Kümülatif sütunlar yılbaşından o aya kadardır. Kapasite = hasar operasyonu ÷ 12 ÷ (0,1 × temel şiddet); kullanım = gelen dosya ÷ kapasite.')
for t in T:
    for m in range(12):
        r = row_of(m, t['id'])
        row(ws, [t['name'], m + 1, r['newPolicies'], r['campaign']['customers'], r['policies'], r['monthGwp'], r['gwp'], r['monthClaims'], r['claimCount'],
                 r['monthLossRatio'], r['grossLossRatio'], r['lossRatio'], r['combinedRatio'], r['capacity'], r['load'], r['utilization'],
                 r['monthService'], r['service'], r['nps'], r['share'], r['monthShare'], r['unitShare'], r['monthProfit'], r['profit'], r['equity']],
            [None, N0, N0, N0, N0, EUR, EUR, EUR, N0, PCT1, PCT1, PCT1, PCT1, N0, N0, PCT, N1, N1, N0, PCT1, PCT1, PCT1, EUR, EUR, EUR])
filt(ws, len(H))

# ——— Gider dökümü ———
H = ['Takım', 'Brüt prim', 'Devredilen prim', 'Net prim', 'Brüt hasar', 'Reasürans payı', 'Net hasar', 'Kanal gideri', 'Sabit gider', 'Karar bütçesi',
     'Reasürans komisyonu', 'Toplam gider', 'Kâr (köprü)', 'Kâr (motor)', 'Fark']
ws = sheet('Gider dökümü', H, [16] + [14] * (len(H) - 1), 'Kâr = net prim − net hasar − (kanal gideri + sabit gider + karar bütçesi − komisyon). Karar bütçesinin tamamı harcanmış sayılır (kampanyanın kullanılmayan hediye parası dahil).')
for t in T:
    r = final[t['id']]
    budget = sum(plan_at(t, m)['marketing'] + plan_at(t, m)['claimsOps'] + (MONEY['reinsuranceFee'] if plan_at(t, m)['reinsurance'] else 0) for m in range(12)) / 12
    commission = r['ceded'] * D['reinsurance']['commission']
    claims = r['netClaims'] + r['recovery']
    fixed = MONEY['fixedMonthly'] * 12
    expenses = r['acquisition'] + fixed + budget - commission
    bridge = r['netEarned'] - r['netClaims'] - expenses
    row(ws, [t['name'], r['gwp'], r['ceded'], r['netEarned'], claims, r['recovery'], r['netClaims'], r['acquisition'], fixed, budget, commission, expenses, bridge, r['profit'], bridge - r['profit']],
        [None] + [EUR] * 13 + [EUR2])

# ——— Kampanya hunisi ———
H = ['Takım', 'Ay', 'Hediye', 'Kampanya payı', 'Medya (ay)', 'Hediye bütçesi (ay)', 'Gösterim', 'Aylık kitle (tüm takımlar)', 'Frekans', 'Erişim', 'Aday (ilgi × tıklama)',
     'İstenen müşteri', 'Kazanılan müşteri', 'Hediye sınırı bağladı mı', 'Kullanılan hediye', 'Boşa giden hediye', 'Müşteri başı maliyet']
ws = sheet('Kampanya hunisi', H, [16, 6, 22] + [13] * (len(H) - 3), 'Aday = erişim × ilgi × tıklama. İstenen = aday × hit (frekans > 5 ise ×1,1) × fiyat etkisi. Kazanılan = min(istenen, hediye bütçesi ÷ hediye maliyeti).')
for t in T:
    for m in range(12):
        k = row_of(m, t['id'])['campaign']
        spend = k['media'] + k['giftBudget']
        row(ws, [t['name'], m + 1, D['offers'][k['offer']]['name'] if k['share'] else '—', k['share'] / 100, k['media'], k['giftBudget'], k['impressions'], k['audience'], k['frequency'], k['reach'], k['leads'],
                 k['wanted'], k['customers'], 'Evet' if k['wanted'] > k['customers'] + 1e-6 else 'Hayır', k['giftsUsed'], k['giftBudget'] - k['giftsUsed'], spend / k['customers'] if k['customers'] > 0 else None],
            [None, N0, None, PCT, EUR, EUR, N0, N0, N2, N0, N0, N1, N1, None, EUR, EUR, EUR2])
filt(ws, len(H))

# ——— Pazar (aylık) ———
H = ['Ay', 'Pazara çıkan müşteri', 'Takımların sattığı (pazar)', 'Kampanya müşterisi', 'Kimseden almayan (piyasanın geri kalanı)', 'Kimseden almayan oranı', 'Takımların brüt primi (ay)', 'Hasar maliyet endeksi', 'Talep endeksi', 'Aktif olaylar']
ws = sheet('Pazar (aylık)', H, [6, 16, 16, 14, 18, 14, 16, 12, 12, 60], 'Pazar havuzunu takımlar ve “piyasanın geri kalanı” paylaşır. Kampanya müşterileri havuzun dışından gelen yeni müşterilerdir.')
for m in range(12):
    mo = MO[m]
    camp = sum(r['campaign']['customers'] for r in mo['rows'])
    sold = sum(r['newPolicies'] for r in mo['rows']) - camp
    row(ws, [m + 1, mo['available'], sold, camp, mo['nonBuyers'], mo['nonBuyers'] / mo['available'], mo['monthGwp'], mo['costIndex'], mo['demandIndex'], ' · '.join(e['title'] for e in mo['events']) or '—'],
        [N0, N0, N0, N0, N0, PCT1, EUR, N2, N2, None])

# ——— Segment sonuçları ———
H = ['Takım', DN['city'], DN['channel'], DN['vehicle'], DN['persona'], DN['type'], 'Pazar poliçesi', 'Kampanya poliçesi', 'Teklif (ilk plan)', 'Referans prim', 'Teklif ÷ referans',
     'Yazılan prim', 'Beklenen hasar', 'Beklenen hasar/prim', 'Hücrenin pazar payı']
ws = sheet('Segment sonuçları', H, [16, 11, 10, 14, 14, 11, 12, 12, 12, 12, 11, 13, 13, 12, 12],
           'Yıllık, takım × hücre. Teklif ilk plandan; prim ve hasar her ayın geçerli planıyla toplanır. Yalnızca en az 1 poliçe satılan hücreler. Olay hasar çarpanları beklenen hasara dahil değildir.')
yearly_pool = [sum(MO[m]['available'] * 0 for m in range(12))]
for t in T:
    for c, cell in enumerate(CELLS):
        mk = sum(next(x for x in MO[m]['trace'] if x['id'] == t['id'])['market'][c] for m in range(12))
        cp = sum(next(x for x in MO[m]['trace'] if x['id'] == t['id'])['campaign'][c] for m in range(12))
        if mk + cp < 1:
            continue
        prem = 0
        for m in range(12):
            st = plan_at(t, m)
            offer = st['basePremium'] * st['coef']['city'][cell['cell'][0]] * st['coef']['channel'][cell['cell'][1]] * st['coef']['vehicle'][cell['cell'][2]] * st['coef']['persona'][cell['cell'][3]] * st['coef']['type'][cell['cell'][4]]
            tr_ = next(x for x in MO[m]['trace'] if x['id'] == t['id'])
            prem += (tr_['market'][c] + tr_['campaign'][c]) * offer
        st0 = t['strategy']
        offer0 = st0['basePremium'] * st0['coef']['city'][cell['cell'][0]] * st0['coef']['channel'][cell['cell'][1]] * st0['coef']['vehicle'][cell['cell'][2]] * st0['coef']['persona'][cell['cell'][3]] * st0['coef']['type'][cell['cell'][4]]
        exp_claims = (mk + cp) * cell['cost']
        yearly = cell['cell'][5] / D['customers'] * D['policies']
        row(ws, [t['name'], *[LV[d][cell['cell'][k]] for k, d in enumerate(['city', 'channel', 'vehicle', 'persona', 'type'])], mk, cp, offer0, cell['reference'], offer0 / cell['reference'],
                 prem, exp_claims, exp_claims / prem if prem else None, mk / yearly if yearly else None],
            [None] * 6 + [N1, N1, EUR2, EUR2, N2, EUR, EUR, PCT1, PCT1])
filt(ws, len(H))

# ——— Final ———
H = ['Sıra (brüt prim)', 'Takım', 'Brüt prim', 'Pazar payı', 'Kâr', 'Brüt hasar/prim', 'NPS', 'Kampanya müşterisi', 'Kârlılık kupası', 'Pay kupası', 'Memnuniyet kupası']
ws = sheet('Final', H, [10, 16, 15, 12, 15, 12, 8, 14, 12, 12, 14])
by = lambda key, tie=None: [r['id'] for r in sorted(final.values(), key=lambda r: (-r[key], -(r[tie] if tie else 0)))]
profit_rank, share_rank, nps_rank = by('profit'), by('share', 'gwp'), by('nps', 'service')
for i, r in enumerate(sorted(final.values(), key=lambda r: -r['gwp']), 1):
    medal = lambda order: {0: '1.', 1: '2.', 2: '3.'}.get(order.index(r['id']), '')
    row(ws, [i, NAME[r['id']], r['gwp'], r['share'], r['profit'], r['grossLossRatio'], r['nps'], r['campaign']['total'], medal(profit_rank), medal(share_rank), medal(nps_rank)],
        [N0, None, EUR, PCT1, EUR, PCT1, N0, N0])

# ——— Mantık kontrolleri ———
ws = sheet('Mantık kontrolleri', ['#', 'Kontrol', 'Beklenen', 'Bulunan', 'Sonuç'], [5, 60, 26, 26, 10])
checks = []
for t in T:
    r = final[t['id']]
    budget = sum(plan_at(t, m)['marketing'] + plan_at(t, m)['claimsOps'] + (MONEY['reinsuranceFee'] if plan_at(t, m)['reinsurance'] else 0) for m in range(12)) / 12
    bridge = r['netEarned'] - r['netClaims'] - (r['acquisition'] + MONEY['fixedMonthly'] * 12 + budget - r['ceded'] * D['reinsurance']['commission'])
    checks.append((f"{t['name']}: kâr kimliği (köprü = motor)", '0 €', f'{bridge - r["profit"]:.6f} €', abs(bridge - r['profit']) < 0.01))
for m in range(12):
    tot = sum(r['share'] for r in MO[m]['rows'])
    checks.append((f'{m + 1}. ay: prim paylarının toplamı', '100%', f'{tot * 100:.6f}%', abs(tot - 1) < 1e-9))
gift_ok = all(row_of(m, t['id'])['campaign']['customers'] * D['offers'][row_of(m, t['id'])['campaign']['offer']]['cost'] <= row_of(m, t['id'])['campaign']['giftBudget'] + 1e-6 for t in T for m in range(12))
checks.append(('Hiçbir ay hediye bütçesinden fazla müşteri kazanılmıyor', 'Evet', 'Evet' if gift_ok else 'Hayır', gift_ok))
sv = D['service']
svc_ok = all(abs(row_of(m, t['id'])['monthService'] - max(sv['floor'], min(sv['max'], sv['max'] - sv['slope'] * max(0, row_of(m, t['id'])['utilization'] - sv['threshold'])))) < 1e-9 for t in T for m in range(12))
checks.append(('Hizmet skoru = 98 − 60 × max(0, kullanım − 0,8), 30–98 arası', 'Her ay', 'Her ay' if svc_ok else 'Sapma var', svc_ok))
nps_ok = all(final[t['id']]['nps'] == round(2 * final[t['id']]['service'] - 100) for t in T)
checks.append(('NPS = 2 × ortalama hizmet − 100', 'Her takım', 'Her takım' if nps_ok else 'Sapma var', nps_ok))
pol_ok = all(abs(sum(row_of(m, t['id'])['newPolicies'] for m in range(12)) - final[t['id']]['policies']) < 1e-6 for t in T)
checks.append(('Aylık yeni poliçeler yıl sonu poliçesine eşit', 'Evet', 'Evet' if pol_ok else 'Hayır', pol_ok))
avail_ok = all(sum(r['newPolicies'] - r['campaign']['customers'] for r in MO[m]['rows']) + MO[m]['nonBuyers'] - MO[m]['available'] < 1e-6 for m in range(12))
checks.append(('Pazar havuzu = takımların sattığı + kimseden almayan', 'Her ay', 'Her ay' if avail_ok else 'Sapma var', avail_ok))
freq_same = all(len({round(r['campaign']['frequency'], 9) for r in MO[m]['rows']}) == 1 for m in range(12))
checks.append(('Kampanya frekansı her ay tüm takımlar için aynı (ortak kitle)', 'Evet', 'Evet' if freq_same else 'Hayır', freq_same))
q_ok = all(row_of(m, T[0]['id'])['campaign']['customers'] == 0 for m in range(6)) and all(row_of(m, T[0]['id'])['campaign']['customers'] > 0 for m in range(6, 12))
checks.append((f"{NAME[T[0]['id']]}: kampanya 7. aydan önce yok, sonra var (çeyrek değişikliği)", 'Evet', 'Evet' if q_ok else 'Hayır', q_ok))
for i, (a, b, c, ok) in enumerate(checks, 1):
    row(ws, [i, a, b, c, 'Geçti' if ok else 'KALDI'], [N0], fill=None)
    ws.cell(ws._next - 1, 5).fill = GOOD if ok else BAD

wb.save('simulasyon-ornegi.xlsx')
print(f"simulasyon-ornegi.xlsx · {len(wb.sheetnames)} sayfa · {sum(1 for x in checks if x[3])}/{len(checks)} kontrol geçti · simulasyon-hesap.xlsx")
